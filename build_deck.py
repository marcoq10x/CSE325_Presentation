#!/usr/bin/env python3
"""
Notespack Frontend - premium animated keynote deck.
Builds a 12-slide widescreen .pptx with sequenced entrance animations
(raw OOXML p:timing) and deliberate slide transitions (p:transition).
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml

# ----------------------------------------------------------------------------
# Design tokens
# ----------------------------------------------------------------------------
BG_DEEP   = "07071A"
BG_MID    = "0E0E22"
BG_CARD   = "141432"
CARD_LINE = "252548"
TEXT      = "EAF1FF"
TEXT_DIM  = "9AA3C4"
TEXT_FAINT= "6E769A"
ACCENT    = "00D4FF"   # cyan
VIOLET    = "6B3EFF"
PINK      = "FF4EA8"

FONT_D = "Calibri"     # display / body, system safe
FONT_B = "Calibri"

EMU_IN = 914400

def C(hex_):
    return RGBColor.from_string(hex_)

# ----------------------------------------------------------------------------
# Presentation scaffold
# ----------------------------------------------------------------------------
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5

# ----------------------------------------------------------------------------
# Low-level helpers
# ----------------------------------------------------------------------------
def add_slide():
    return prs.slides.add_slide(BLANK)

def _spPr(shape):
    return shape._element.spPr

def _clear_fill(spPr):
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)

def _insert_fill(spPr, fill_el):
    _clear_fill(spPr)
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addprevious(fill_el)
    else:
        spPr.append(fill_el)

def solid(shape, hex_, alpha=None):
    a = "" if alpha is None else f'<a:alpha val="{int(alpha*1000)}"/>'
    xml = f'<a:solidFill {nsdecls("a")}><a:srgbClr val="{hex_}">{a}</a:srgbClr></a:solidFill>'
    _insert_fill(_spPr(shape), parse_xml(xml))

def no_fill(shape):
    _insert_fill(_spPr(shape), parse_xml(f'<a:noFill {nsdecls("a")}/>'))

def gradient(shape, stops, angle=135, radial=False):
    """stops: list of (pos_float_0_1, hex, alpha_percent_or_None)."""
    gs = []
    for stop in stops:
        pos, hx = stop[0], stop[1]
        al = stop[2] if len(stop) > 2 else None
        a = "" if al is None else f'<a:alpha val="{int(al*1000)}"/>'
        gs.append(f'<a:gs pos="{int(pos*100000)}"><a:srgbClr val="{hx}">{a}</a:srgbClr></a:gs>')
    if radial:
        geom = '<a:path path="circle"><a:fillToRect l="100000" t="100000" r="0" b="0"/></a:path>'
    else:
        geom = f'<a:lin ang="{int(angle*60000)}" scaled="1"/>'
    xml = (f'<a:gradFill {nsdecls("a")} rotWithShape="1"><a:gsLst>'
           + "".join(gs) + f'</a:gsLst>{geom}</a:gradFill>')
    _insert_fill(_spPr(shape), parse_xml(xml))

def line(shape, hex_=None, width_pt=None, alpha=None):
    if hex_ is None:
        shape.line.fill.background()
        return
    shape.line.color.rgb = C(hex_)
    if width_pt is not None:
        shape.line.width = Pt(width_pt)
    if alpha is not None:
        srgb = shape.line.color._xFill.find(qn("a:srgbClr"))
        srgb.append(parse_xml(f'<a:alpha {nsdecls("a")} val="{int(alpha*1000)}"/>'))

def no_line(shape):
    shape.line.fill.background()

def soft_shadow(shape, blur=10, dist=6, dir=5400000, alpha=55):
    """Subtle outer drop shadow for depth."""
    spPr = _spPr(shape)
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    xml = (f'<a:effectLst {nsdecls("a")}>'
           f'<a:outerShdw blurRad="{int(blur*12700)}" dist="{int(dist*12700)}" '
           f'dir="{dir}" rotWithShape="0"><a:srgbClr val="000000">'
           f'<a:alpha val="{int(alpha*1000)}"/></a:srgbClr></a:outerShdw></a:effectLst>')
    spPr.append(parse_xml(xml))

def glow(shape, hex_, rad=14, alpha=60):
    spPr = _spPr(shape)
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    xml = (f'<a:effectLst {nsdecls("a")}><a:glow rad="{int(rad*12700)}">'
           f'<a:srgbClr val="{hex_}"><a:alpha val="{int(alpha*1000)}"/></a:srgbClr>'
           f'</a:glow></a:effectLst>')
    spPr.append(parse_xml(xml))

def rect(slide, x, y, w, h, rounded=False, radius=0.08):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp

def oval(slide, x, y, w, h):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    return shp

def rotate(shape, deg):
    shape.rotation = deg

def _apply_run(run, size, hex_, bold, font, spc, alpha):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "en-US")
    # color (with optional alpha)
    for sf in rPr.findall(qn("a:solidFill")):
        rPr.remove(sf)
    a = "" if alpha is None else f'<a:alpha val="{int(alpha*1000)}"/>'
    rPr.insert(0, parse_xml(f'<a:solidFill {nsdecls("a")}><a:srgbClr val="{hex_}">{a}</a:srgbClr></a:solidFill>'))
    if spc is not None:
        rPr.set("spc", str(int(spc * 100)))
    # ensure latin/cs font
    for tag in ("a:latin", "a:cs"):
        for el in rPr.findall(qn(tag)):
            rPr.remove(el)
    rPr.append(parse_xml(f'<a:latin {nsdecls("a")} typeface="{font}"/>'))

def text(slide, x, y, w, h, content, size=14, color=TEXT, bold=False,
         font=FONT_D, align=PP_ALIGN.LEFT, caps=False, spc=None, alpha=None,
         anchor=MSO_ANCHOR.TOP, line_spacing=None, wrap=True):
    """content: str (single para) or list of (text, overrides-dict) tuples for one para,
       or list of paragraphs where each paragraph is str or list-of-runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0

    paragraphs = content if isinstance(content, list) and content and isinstance(content[0], (list, str)) and not (isinstance(content[0], tuple)) else [content]
    # normalize: we treat list of str/lists as multiple paragraphs; single str => one para
    if isinstance(content, str):
        paragraphs = [content]
    elif isinstance(content, list):
        paragraphs = content

    first = True
    for para in paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        runs = para if isinstance(para, list) else [(para, {})]
        for rt, ov in runs:
            r = p.add_run()
            txt = rt.upper() if caps else rt
            r.text = txt
            _apply_run(r, ov.get("size", size), ov.get("color", color),
                       ov.get("bold", bold), ov.get("font", font),
                       ov.get("spc", spc), ov.get("alpha", alpha))
    return tb

# ----------------------------------------------------------------------------
# Recurring components
# ----------------------------------------------------------------------------
def bg(slide, hex_=BG_DEEP):
    r = rect(slide, -0.06, -0.06, SW + 0.12, SH + 0.12)
    solid(r, hex_); no_line(r)
    return r

def blob(slide, x, y, w, h, stops, angle, rot, alpha_each):
    r = rect(slide, x, y, w, h, rounded=True, radius=0.5)
    gradient(r, [(p, hx, alpha_each) for (p, hx) in stops], angle=angle)
    no_line(r); rotate(r, rot)
    return r

def footer(slide, number, anim=None):
    ln = rect(slide, 0.6, 7.1, SW - 1.2, 0.012)
    solid(ln, ACCENT, alpha=55); no_line(ln)
    text(slide, 0.6, 7.14, 6, 0.3, "NOTESPACK · FRONTEND BUILD",
         size=9, color=TEXT_FAINT, bold=True, spc=2.2, caps=True)
    text(slide, SW - 1.6, 7.14, 1.0, 0.3, f"{number:02d}",
         size=9, color=TEXT_FAINT, bold=True, align=PP_ALIGN.RIGHT)
    return ln

def eyebrow(slide, label, x=0.6, y=0.6, color=ACCENT):
    return text(slide, x, y, 9, 0.35, label, size=11, color=color, bold=True,
                spc=3.0, caps=True)

def card(slide, x, y, w, h, fill=BG_CARD, line_hex=CARD_LINE, line_w=0.75,
         radius=0.06, shadow=True):
    c = rect(slide, x, y, w, h, rounded=True, radius=radius)
    solid(c, fill); line(c, line_hex, line_w)
    if shadow:
        soft_shadow(c, blur=18, dist=7, alpha=42)
    return c

# ----------------------------------------------------------------------------
# Animation engine (raw OOXML p:timing) + transitions
# ----------------------------------------------------------------------------
class Anim:
    """Collects entrance animations for a slide and emits auto-playing timing XML."""
    def __init__(self):
        self.items = []   # (spid, effect, delay_ms, dur_ms)
    def add(self, shape, effect="rise", delay=0.0, dur=0.55):
        self.items.append((shape.shape_id, effect, int(delay * 1000), int(dur * 1000)))
        return shape
    def group(self, shapes, effect="rise", start=0.0, stagger=0.1, dur=0.45):
        for i, s in enumerate(shapes):
            self.add(s, effect, start + i * stagger, dur)

P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"

def _effect_children(spid, effect, dur_ms, cid):
    """Return (xml_string, next_cid) for the inner behaviors of one effect."""
    rise_y = 0.022  # ~12px upward on a 7.5in slide
    setvis = (f'<p:set><p:cBhvr>'
              f'<p:cTn id="{cid}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
              f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
              f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
              f'</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>')
    cid += 1
    if effect == "wipe":
        fade = (f'<p:animEffect transition="in" filter="wipe(right)">'
                f'<p:cBhvr><p:cTn id="{cid}" dur="{dur_ms}"/>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')
        cid += 1
        return setvis + fade, cid
    if effect == "wipe_down":
        fade = (f'<p:animEffect transition="in" filter="wipe(down)">'
                f'<p:cBhvr><p:cTn id="{cid}" dur="{dur_ms}"/>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')
        cid += 1
        return setvis + fade, cid
    # fade (base of fade + rise)
    fade = (f'<p:animEffect transition="in" filter="fade">'
            f'<p:cBhvr><p:cTn id="{cid}" dur="{dur_ms}"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')
    cid += 1
    if effect == "fade":
        return setvis + fade, cid
    # rise = fade + short upward motion
    motion = (f'<p:animMotion origin="layout" path="M 0 {rise_y} L 0 0" '
              f'pathEditMode="relative" rAng="0" ptsTypes="">'
              f'<p:cBhvr><p:cTn id="{cid}" dur="{dur_ms}" fill="hold"/>'
              f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
              f'<p:attrNameLst><p:attrName>ppt_x</p:attrName>'
              f'<p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>')
    cid += 1
    return setvis + fade + motion, cid

PRESET = {
    "fade": (10, 0), "rise": (10, 0),
    "wipe": (22, 2), "wipe_down": (22, 8),
}

def build_timing(anim):
    items = sorted(anim.items, key=lambda t: t[2])
    cid = 5  # ids 1..4 used by scaffold
    pars = []
    for spid, effect, delay_ms, dur_ms in items:
        pid = cid; cid += 1
        children, cid = _effect_children(spid, effect, dur_ms, cid)
        presetID, subtype = PRESET[effect]
        node = "withEffect"
        pars.append(
            f'<p:par><p:cTn id="{pid}" presetID="{presetID}" presetClass="entr" '
            f'presetSubtype="{subtype}" fill="hold" grpId="0" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="{delay_ms}"/></p:stCondLst>'
            f'<p:childTnLst>{children}</p:childTnLst></p:cTn></p:par>'
        )
    inner = "".join(pars)
    xml = (
        f'<p:timing {nsdecls("p")}>'
        f'<p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
        f'<p:childTnLst><p:seq concurrent="1" nextAc="seek">'
        f'<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        f'<p:childTnLst><p:par><p:cTn id="4" fill="hold">'
        f'<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        f'<p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        f'</p:par></p:childTnLst></p:cTn>'
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
    )
    return parse_xml(xml)

def apply_anim(slide, anim):
    sld = slide._element
    for old in sld.findall(qn("p:timing")):
        sld.remove(old)
    sld.append(build_timing(anim))

def transition(slide, kind="push", dur=700):
    sld = slide._element
    for old in sld.findall(qn("p:transition")):
        sld.remove(old)
    if kind == "morph":
        inner = '<p14:morph option="byObject"/>'
    elif kind == "fadeblk":
        inner = '<p:fade thruBlk="1"/>'
    elif kind == "push":
        inner = '<p:push dir="r"/>'
    else:
        inner = '<p:fade/>'
    MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    fb = inner if kind != "morph" else "<p:fade/>"
    xml = (f'<mc:AlternateContent xmlns:mc="{MC}" {nsdecls("p")} xmlns:p14="{P14}">'
           f'<mc:Choice Requires="p14">'
           f'<p:transition spd="med" p14:dur="{dur}">{inner}</p:transition>'
           f'</mc:Choice><mc:Fallback>'
           f'<p:transition spd="med">{fb}</p:transition>'
           f'</mc:Fallback></mc:AlternateContent>')
    el = parse_xml(xml)
    cSld = sld.find(qn("p:cSld"))
    clrMap = sld.find(qn("p:clrMapOvr"))
    ref = clrMap if clrMap is not None else cSld
    ref.addnext(el)

# ============================================================================
# SLIDE 1 - COVER
# ============================================================================
def slide1():
    s = add_slide(); a = Anim()
    bg(s, BG_DEEP)
    b1 = blob(s, 8.0, -3.0, 12, 8, [(0, ACCENT), (1, VIOLET)], 135, 25, 18)
    b2 = blob(s, -2.0, 4.0, 9, 6, [(0, VIOLET), (1, PINK)], 135, -15, 12)

    eb = eyebrow(s, "FRONTEND BUILD · CSE 325")
    title = text(s, 0.55, 1.55, 12, 2.0, "Notespack", size=128, color=TEXT, bold=True)
    tag = text(s, 0.6, 3.55, 11, 0.7, "A campus events platform, built for students.",
               size=24, color=TEXT_DIM)
    rule = rect(s, 0.6, 4.5, 4.8, 0.055, rounded=True, radius=0.5)
    gradient(rule, [(0, ACCENT), (1, VIOLET)], angle=0); no_line(rule)

    meta = [("PRESENTER", "Marco"), ("TEAM", "The Syntax Errors"),
            ("COURSE", "CSE 325 · BYU-Idaho"), ("BRANCH", "marco/frontend-design")]
    cells = []
    cw = (SW - 1.2) / 4
    for i, (lab, val) in enumerate(meta):
        cx = 0.6 + i * cw
        if i > 0:
            d = rect(s, cx - 0.05, 6.05, 0.008, 0.85); solid(d, ACCENT, alpha=30); no_line(d)
        l = text(s, cx, 6.0, cw - 0.3, 0.3, lab, size=9, color=TEXT_FAINT, bold=True, spc=2.5, caps=True)
        v = text(s, cx, 6.35, cw - 0.2, 0.5, val, size=14, color=TEXT, bold=True)
        cells.append((l, v))

    a.add(b1, "fade", 0.0, 0.8); a.add(b2, "fade", 0.1, 0.8)
    a.add(eb, "rise", 0.4, 0.5)
    a.add(title, "rise", 0.6, 0.7)
    a.add(tag, "rise", 1.0, 0.5)
    a.add(rule, "wipe", 1.3, 0.6)
    flat = [x for c in cells for x in c]
    a.group(flat, "rise", start=1.6, stagger=0.06, dur=0.4)
    apply_anim(s, a)
    return s

# ============================================================================
# SLIDE 2 - WHAT IS NOTESPACK
# ============================================================================
def slide2():
    s = add_slide(); a = Anim()
    bg(s, BG_DEEP)
    blob(s, 9.5, -2.5, 7, 6, [(0, ACCENT), (1, VIOLET)], 135, 20, 9)
    eb = eyebrow(s, "01 · CONTEXT")
    title = text(s, 0.6, 1.0, 11, 0.9, "What is Notespack?", size=44, color=TEXT, bold=True)

    para = text(s, 0.6, 2.45, 7.0, 3.5,
                "Notespack helps students discover, share, and attend events around campus. "
                "Browse what's happening, filter by what you care about, and create your own "
                "events to invite friends. Built as a group project for CSE 325.",
                size=18, color=TEXT_DIM, line_spacing=1.5)

    cd = card(s, 8.2, 2.45, 4.5, 3.6)
    rows = [("WHO", "Campus students"), ("WHAT", "Events platform"),
            ("MY PART", "Entire frontend homepage")]
    row_shapes = []
    ry = 2.8
    for i, (lab, val) in enumerate(rows):
        l = text(s, 8.5, ry, 4.0, 0.3, lab, size=9, color=ACCENT, bold=True, spc=2.0, caps=True)
        v = text(s, 8.5, ry + 0.28, 4.0, 0.4, val, size=14, color=TEXT, bold=True)
        row_shapes += [l, v]
        if i < 2:
            dv = rect(s, 8.5, ry + 0.82, 3.9, 0.008); solid(dv, CARD_LINE); no_line(dv)
        ry += 1.05

    footer(s, 2)
    a.add(eb, "rise", 0.0, 0.5); a.add(title, "rise", 0.05, 0.5)
    a.add(para, "rise", 0.3, 0.6)
    a.add(cd, "rise", 0.5, 0.5)
    a.group(row_shapes, "rise", start=0.8, stagger=0.1, dur=0.4)
    apply_anim(s, a)
    transition(s, "morph", 800)
    return s

# ============================================================================
# SLIDE 3 - WHAT I SHIPPED (stat hero)
# ============================================================================
def slide3():
    s = add_slide(); a = Anim()
    bg(s, BG_DEEP)
    eb = eyebrow(s, "02 · SCOPE")
    title = text(s, 0.6, 1.0, 11, 0.9, "What I shipped", size=44, color=TEXT, bold=True)

    cols = [
        ("9", ACCENT, "HOMEPAGE SECTIONS",
         "Hero · Marquee · Categories · Why · Stats · Featured · Testimonials · How · CTA"),
        ("6", VIOLET, "FILES CHANGED",
         "Index.cshtml · _Layout.cshtml · Privacy · site.css · site.js · NPlogo.png"),
        ("0", PINK, "BACKEND FILES TOUCHED",
         "Rosana's models and controllers are untouched. Pure frontend."),
    ]
    xs = [0.6, 4.7, 8.8]; w = 4.0
    # dividers
    for dx in (4.55, 8.65):
        d = rect(s, dx, 2.8, 0.008, 3.8); solid(d, CARD_LINE); no_line(d)
    seq = []
    for i, (num, col, lab, cap) in enumerate(cols):
        x = xs[i]
        n = text(s, x, 2.5, w, 2.2, num, size=200, color=col, bold=True, align=PP_ALIGN.CENTER)
        ln = rect(s, x + w/2 - 0.75, 4.95, 1.5, 0.016, rounded=True, radius=0.5)
        solid(ln, col, alpha=45); no_line(ln)
        l = text(s, x, 5.15, w, 0.4, lab, size=11, color=TEXT_FAINT, bold=True, spc=2.0, caps=True, align=PP_ALIGN.CENTER)
        c = text(s, x + 0.2, 5.55, w - 0.4, 1.0, cap, size=10, color=TEXT_DIM, align=PP_ALIGN.CENTER, line_spacing=1.2)
        seq.append((n, ln, l, c))

    footer(s, 3)
    a.add(eb, "rise", 0.0, 0.5); a.add(title, "rise", 0.05, 0.5)
    for i, (n, ln, l, c) in enumerate(seq):
        base = 0.4 + i * 0.4
        a.add(n, "rise", base, 0.6)
        a.add(ln, "wipe", base + 0.3, 0.4)
        a.add(l, "fade", base + 0.5, 0.4)
        a.add(c, "fade", base + 0.55, 0.4)
    apply_anim(s, a)
    transition(s, "push", 700)
    return s

# ============================================================================
# SLIDE 4 - NINE SECTIONS
# ============================================================================
def slide4():
    s = add_slide(); a = Anim()
    bg(s, BG_DEEP)
    eb = eyebrow(s, "03 · HOMEPAGE TOUR")
    title = text(s, 0.6, 1.0, 11, 0.9, "Nine sections, top to bottom", size=44, color=TEXT, bold=True)

    spine = rect(s, 1.05, 2.2, 0.08, 4.75, rounded=True, radius=0.5)
    gradient(spine, [(0, ACCENT), (0.5, VIOLET), (1, PINK)], angle=90); no_line(spine)

    sections = [
        ("Hero", "Animated split text, gradient rule, dual call-to-action buttons"),
        ("Event marquee", "Glowing horizontal scroll of event names"),
        ("Categories ribbon", "Pills with active state, ready for filter wiring"),
        ("Why Notespack", "Three glass cards with 3D tilt and cursor glow"),
        ("Stat counters", "Numbers animate from zero when scrolled into view"),
        ("Featured events", "Six cards with gradient covers and date badges"),
        ("Testimonials", "Student quotes with avatar gradient circles"),
        ("How it works", "Three numbered steps from sign-in to attend"),
        ("CTA strip", "Secondary conversion path before the footer"),
    ]
    node_colors = [ACCENT]*3 + [VIOLET]*3 + [PINK]*3
    nodes = []
    for i, (name, desc) in enumerate(sections):
        cy = 2.2 + i * 0.525
        circ = oval(s, 0.86, cy, 0.34, 0.34)
        solid(circ, node_colors[i]); no_line(circ)
        glow(circ, node_colors[i], rad=8, alpha=45)
        num = text(s, 0.86, cy + 0.02, 0.34, 0.3, f"{i+1:02d}", size=9, color="07071A", bold=True, align=PP_ALIGN.CENTER)
        nm = text(s, 1.45, cy - 0.05, 6.7, 0.3, name, size=15, color=TEXT, bold=True)
        ds = text(s, 1.45, cy + 0.23, 6.7, 0.3, desc, size=11, color=TEXT_DIM)
        nodes.append((circ, num, nm, ds))

    pc = card(s, 8.5, 2.45, 4.3, 4.4)
    text(s, 8.8, 2.7, 3.7, 0.3, "PLUS", size=11, color=ACCENT, bold=True, spc=3.0, caps=True)
    plus = [
        "Circular chrome NP navbar logo",
        "Skip-to-content link for keyboard users",
        "Multi-column footer (Brand · Product · For Students · Company)",
    ]
    py = 3.25
    plus_shapes = []
    for i, item in enumerate(plus):
        dot = oval(s, 8.8, py + 0.06, 0.12, 0.12); solid(dot, [ACCENT, VIOLET, PINK][i]); no_line(dot)
        t = text(s, 9.1, py, 3.4, 0.7, item, size=13, color=TEXT_DIM, line_spacing=1.15)
        plus_shapes += [dot, t]
        py += 1.05

    footer(s, 4)
    a.add(eb, "rise", 0.0, 0.5); a.add(title, "rise", 0.05, 0.5)
    a.add(spine, "wipe_down", 0.4, 0.8)
    flat = [x for n in nodes for x in n]
    # stagger by node group
    for i, grp in enumerate(nodes):
        st = 0.9 + i * 0.08
        for sh in grp:
            a.add(sh, "rise", st, 0.35)
    a.add(pc, "rise", 1.7, 0.5)
    a.group(plus_shapes, "rise", start=1.9, stagger=0.08, dur=0.4)
    apply_anim(s, a)
    transition(s, "push", 700)
    return s

# ============================================================================
# SLIDE 5 - LITTLE TOUCHES (3x2 cards)
# ============================================================================
def slide5():
    s = add_slide(); a = Anim()
    bg(s, BG_DEEP)
    eb = eyebrow(s, "04 · INTERACTIONS")
    title = text(s, 0.6, 1.0, 11, 0.9, "The little touches, in plain words", size=44, color=TEXT, bold=True)
    sub = text(s, 0.6, 1.75, 12, 0.4,
               "Just vanilla JavaScript plus a tiny GSAP cameo for the hero entrance. No heavy libraries.",
               size=14, color=TEXT_DIM)

    cards = [
        ("3D card tilt", "G1", "Cards lean toward your mouse as you move across them. Subtle, not loud."),
        ("Cursor glow", "G2", "Each card has a soft cyan light inside that follows your cursor."),
        ("Cursor halo", "G1", "A faint cyan dot trails your mouse everywhere on the page."),
        ("Smooth fade-ins", "G2", "Sections gently fade up as you scroll down. No sudden jumps."),
        ("Counting numbers", "G1", "The stat numbers count from zero to their target when you scroll past."),
        ("Active pill", "G2", "Clicked category pills fill with cyan to show what's selected."),
    ]
    cw, ch = 4.0, 2.05
    gx, gy = 0.22, 0.2
    x0, y0 = 0.6, 2.35
    order_shapes = []
    for i, (ttl, grad, desc) in enumerate(cards):
        col = i % 3; row = i // 3
        x = x0 + col * (cw + gx); y = y0 + row * (ch + gy)
        cd = card(s, x, y, cw, ch)
        chip = rect(s, x + 0.3, y + 0.3, 0.42, 0.42, rounded=True, radius=0.28)
        if grad == "G1":
            gradient(chip, [(0, ACCENT), (1, VIOLET)], 135)
        else:
            gradient(chip, [(0, VIOLET), (1, PINK)], 135)
        no_line(chip); glow(chip, ACCENT if grad=="G1" else PINK, rad=6, alpha=40)
        h = text(s, x + 0.95, y + 0.32, cw - 1.2, 0.4, ttl, size=16, color=TEXT, bold=True)
        d = text(s, x + 0.3, y + 0.95, cw - 0.6, 1.0, desc, size=13, color=TEXT_DIM, line_spacing=1.25)
        order_shapes.append([cd, chip, h, d])

    footer(s, 5)
    a.add(eb, "rise", 0.0, 0.5); a.add(title, "rise", 0.05, 0.5)
    a.add(sub, "rise", 0.2, 0.5)
    for i, grp in enumerate(order_shapes):
        st = 0.6 + i * 0.1
        for sh in grp:
            a.add(sh, "rise", st, 0.45)
    apply_anim(s, a)
    transition(s, "push", 700)
    return s

# ============================================================================
# SLIDE 6 - ACCESSIBILITY (full-bleed G1)
# ============================================================================
def slide6():
    s = add_slide(); a = Anim()
    base = rect(s, -0.06, -0.06, SW + 0.12, SH + 0.12)
    gradient(base, [(0, ACCENT), (1, VIOLET)], 135); no_line(base)

    eb = eyebrow(s, "05 · ACCESSIBILITY", color="FFFFFF")
    head = text(s, 0.6, 1.9, 12.1, 1.3, "a11y means accessibility.", size=72, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
    subx = text(s, 0.6, 3.25, 12.1, 0.7, "The 11 stands for the eleven letters between 'a' and 'y'.",
                size=28, color="FFFFFF", align=PP_ALIGN.CENTER, alpha=85)

    examples = [
        ("Blind users", "Navigate using screen readers"),
        ("Keyboard only", "Tab key, arrow keys, no mouse"),
        ("Low vision", "Need high contrast and big text"),
        ("Color blind", "Cannot rely on color alone"),
    ]
    cw = (SW - 1.2) / 4
    ex_shapes = []
    for i, (lab, desc) in enumerate(examples):
        cx = 0.6 + i * cw
        circ = oval(s, cx + cw/2 - 0.18, 4.7, 0.36, 0.36)
        no_fill(circ); line(circ, "FFFFFF", 1.25)
        l = text(s, cx, 5.2, cw, 0.35, lab, size=14, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
        d = text(s, cx, 5.55, cw, 0.6, desc, size=12, color="FFFFFF", align=PP_ALIGN.CENTER, alpha=82, line_spacing=1.1)
        ex_shapes.append([circ, l, d])
    micro = text(s, 0.6, 6.55, 12.1, 0.4, "Our homepage works for all of them.",
                 size=16, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)

    a.add(eb, "fade", 0.0, 0.4)
    a.add(head, "rise", 0.4, 0.7)
    a.add(subx, "rise", 1.0, 0.5)
    for i, grp in enumerate(ex_shapes):
        st = 1.5 + i * 0.12
        for sh in grp: a.add(sh, "rise", st, 0.4)
    a.add(micro, "fade", 2.2, 0.4)
    apply_anim(s, a)
    transition(s, "fadeblk", 800)
    return s

# ============================================================================
# SLIDE 7 - WCAG (full-bleed G2)
# ============================================================================
def slide7():
    s = add_slide(); a = Anim()
    base = rect(s, -0.06, -0.06, SW + 0.12, SH + 0.12)
    gradient(base, [(0, VIOLET), (1, PINK)], 135); no_line(base)

    eb = eyebrow(s, "06 · STANDARDS", color="FFFFFF")
    head = text(s, 0.6, 1.7, 12.1, 1.1, "WCAG 2.1 AA", size=56, color="FFFFFF", bold=True)
    subx = text(s, 0.6, 2.85, 12.1, 0.6, "Web Content Accessibility Guidelines, version 2.1, level AA.",
                size=24, color="FFFFFF", alpha=85)

    cd = rect(s, 2.5, 3.95, 8.3, 2.55, rounded=True, radius=0.05)
    solid(cd, "FFFFFF", alpha=12); line(cd, "FFFFFF", 1.0, alpha=30)

    lh = text(s, 2.85, 4.25, 3.5, 0.3, "WHAT IT IS", size=11, color="FFFFFF", bold=True, spc=2.0, caps=True)
    lb = text(s, 2.85, 4.6, 3.5, 1.8,
              "The global rulebook for web accessibility. Three levels: A (minimum), AA "
              "(industry standard), AAA (strictest). AA is what BYU, government sites, and most "
              "companies require.", size=13, color="FFFFFF", line_spacing=1.2, alpha=92)
    rh = text(s, 6.7, 4.25, 3.8, 0.3, "WHAT WE PASS", size=11, color="FFFFFF", bold=True, spc=2.0, caps=True)
    checks = ["Text contrast 4.5 to 1", "Keyboard navigation works",
              "Focus rings visible", "Reduced motion respected"]
    chk_shapes = []
    cy = 4.62
    for c in checks:
        t = text(s, 6.7, cy, 3.9, 0.35, [[("✓  ", {"color":"FFFFFF","bold":True}), (c, {})]],
                 size=13, color="FFFFFF", alpha=95)
        chk_shapes.append(t); cy += 0.42
    micro = text(s, 0.6, 6.75, 12.1, 0.4, "The homepage passes.", size=16, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)

    a.add(eb, "fade", 0.0, 0.4)
    a.add(head, "rise", 0.3, 0.7)
    a.add(subx, "fade", 0.9, 0.5)
    a.add(cd, "rise", 1.3, 0.5)
    a.add(lh, "fade", 1.7, 0.4); a.add(lb, "fade", 1.75, 0.4)
    a.add(rh, "fade", 1.9, 0.3)
    a.group(chk_shapes, "rise", start=2.0, stagger=0.1, dur=0.3)
    a.add(micro, "fade", 2.7, 0.4)
    apply_anim(s, a)
    transition(s, "fadeblk", 800)
    return s

# ============================================================================
# SLIDE 8 - RESPONSIVE
# ============================================================================
def slide8():
    s = add_slide(); a = Anim()
    bg(s, BG_DEEP)
    eb = eyebrow(s, "07 · RESPONSIVE")
    title = text(s, 0.6, 1.0, 11, 0.9, "Works on every screen size", size=44, color=TEXT, bold=True)
    sub = text(s, 0.6, 1.7, 12, 0.4,
               "Four breakpoints. The layout adapts so the page never feels cramped or stretched.",
               size=14, color=TEXT_DIM)

    # device: (x, y, w, h, color, name, range, behavior)
    devices = [
        (0.7, 2.55, 4.0, 2.45, ACCENT, "Desktop", "≥ 900px", "Full 4-column footer, full nav"),
        (5.35, 2.65, 2.6, 2.25, VIOLET, "Tablet", "720-900px", "Two-column footer"),
        (8.35, 2.55, 1.5, 2.55, PINK, "Mobile", "480-720px", "Stacked layout, condensed nav"),
        (10.6, 2.8, 1.0, 2.05, PINK, "Phone", "< 480px", "Logo only, wordmark hidden"),
    ]
    LBOX = 2.0
    dev_groups = []
    for (x, y, w, h, col, name, rng, beh) in devices:
        cx = x + w/2
        frame = rect(s, x, y, w, h, rounded=True, radius=0.08)
        solid(frame, BG_CARD); line(frame, col, 1.5)
        glow(frame, col, rad=8, alpha=30)
        blocks = []
        bx = x + 0.18; bw = w - 0.36
        by = y + 0.28
        bh = 0.16
        for k in range(3):
            blk = rect(s, bx, by, bw, bh, rounded=True, radius=0.4)
            solid(blk, col, alpha=38); no_line(blk)
            blocks.append(blk)
            by += bh + 0.16
        nm = text(s, cx - LBOX/2, 5.4, LBOX, 0.3, name, size=14, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        rg = text(s, cx - LBOX/2, 5.73, LBOX, 0.3, rng, size=11, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
        bv = text(s, cx - LBOX/2, 6.05, LBOX, 0.8, beh, size=11, color=TEXT_DIM, align=PP_ALIGN.CENTER, line_spacing=1.1)
        dev_groups.append((frame, blocks, nm, rg, bv))

    footer(s, 8)
    a.add(eb, "rise", 0.0, 0.5); a.add(title, "rise", 0.05, 0.5)
    a.add(sub, "rise", 0.2, 0.5)
    for i, (frame, blocks, nm, rg, bv) in enumerate(dev_groups):
        st = 0.5 + i * 0.15
        a.add(frame, "rise", st, 0.5)
        for j, blk in enumerate(blocks):
            a.add(blk, "fade", st + 0.05 + j*0.04, 0.3)
        a.add(nm, "fade", 1.3 + i*0.1, 0.3)
        a.add(rg, "fade", 1.35 + i*0.1, 0.3)
        a.add(bv, "fade", 1.4 + i*0.1, 0.3)
    apply_anim(s, a)
    transition(s, "push", 700)
    return s

# ============================================================================
# SLIDE 9 - ASKED VS SHIPPED
# ============================================================================
def slide9():
    s = add_slide(); a = Anim()
    bg(s, BG_DEEP)
    eb = eyebrow(s, "08 · PROPOSAL COVERAGE")
    title = text(s, 0.6, 1.0, 12, 0.9, "What was asked, what I shipped", size=44, color=TEXT, bold=True)

    cols = [
        (ACCENT, "DONE · 4", "✓",
         ["Browse events (Featured Events grid)", "See categories (ribbon)",
          "Mobile responsive", "Modern UI"]),
        (VIOLET, "UI READY · 1", "•",
         ["Filter by category. Pills wired to click, waiting on the filter endpoint."]),
        (PINK, "PENDING · 5", "○",
         ["Sign in / Register", "Create event", "Edit / Delete event",
          "Share event with code", "Search"]),
    ]
    cw = 4.0; gx = 0.15; x0 = 0.6; y0 = 2.4; ch = 4.4
    col_groups = []
    for i, (col, hdr, glyph, items) in enumerate(cols):
        x = x0 + i * (cw + gx)
        cd = card(s, x, y0, cw, ch)
        hb = rect(s, x, y0, cw, 0.5, rounded=True, radius=0.12)
        solid(hb, col); no_line(hb)
        # square off the bottom of the header bar by overlaying
        hb2 = rect(s, x, y0 + 0.25, cw, 0.25); solid(hb2, col); no_line(hb2)
        ht = text(s, x, y0 + 0.08, cw, 0.35, hdr, size=12, color=BG_DEEP, bold=True, spc=1.5, caps=True, align=PP_ALIGN.CENTER)
        item_shapes = []
        iy = y0 + 0.75
        for it in items:
            t = text(s, x + 0.3, iy, cw - 0.6, 0.7,
                     [[(glyph + "  ", {"color": col, "bold": True}), (it, {})]],
                     size=13, color=TEXT, line_spacing=1.15)
            item_shapes.append(t)
            iy += 0.66 if len(it) < 36 else 0.86
        col_groups.append((cd, hb, hb2, ht, item_shapes))

    footer(s, 9)
    a.add(eb, "rise", 0.0, 0.5); a.add(title, "rise", 0.05, 0.5)
    for i, (cd, hb, hb2, ht, items) in enumerate(col_groups):
        st = 0.5 + i * 0.2
        a.add(cd, "rise", st, 0.5)
        a.add(hb, "fade", st + 0.05, 0.4); a.add(hb2, "fade", st + 0.05, 0.4)
        a.add(ht, "fade", st + 0.1, 0.4)
        for j, it in enumerate(items):
            a.add(it, "rise", st + 0.3 + j * 0.1, 0.3)
    apply_anim(s, a)
    transition(s, "push", 700)
    return s

# ============================================================================
# SLIDE 10 - SIX FILES
# ============================================================================
def slide10():
    s = add_slide(); a = Anim()
    bg(s, BG_DEEP)
    eb = eyebrow(s, "09 · CHANGE LOG")
    title = text(s, 0.6, 1.0, 12, 0.9, "The six files I changed", size=44, color=TEXT, bold=True)
    sub = text(s, 0.6, 1.7, 12, 0.4, "Pure frontend. Rosana's backend scaffold is untouched.",
               size=14, color=TEXT_DIM)

    files = [
        ("Views/Home/Index.cshtml", "9 homepage sections, all the magic happens here"),
        ("Views/Shared/_Layout.cshtml", "Circular navbar logo, 4-column footer"),
        ("Views/Home/Privacy.cshtml", "Themed privacy page to match"),
        ("wwwroot/css/site.css", "Full design system, breakpoints, motion"),
        ("wwwroot/js/site.js", "Reveals, counters, cursor glow, 3D tilt"),
        ("wwwroot/img/NPlogo.png", "Chrome NP brand asset"),
    ]
    rows = []
    ry = 2.45
    for i, (path, purpose) in enumerate(files):
        bar = rect(s, 0.6, ry, 12.1, 0.5, rounded=True, radius=0.12)
        solid(bar, BG_CARD, alpha=70); no_line(bar)
        dot = rect(s, 0.6, ry, 0.06, 0.5, rounded=True, radius=0.5)
        gradient(dot, [(0, ACCENT), (1, VIOLET)], 90); no_line(dot)
        fp = text(s, 0.95, ry + 0.1, 6.5, 0.35, path, size=13, color=TEXT, bold=True)
        pр = text(s, 7.5, ry + 0.12, 5.1, 0.35, purpose, size=12, color=TEXT_DIM, align=PP_ALIGN.RIGHT)
        rows.append([bar, dot, fp, pр])
        ry += 0.62

    pill = rect(s, 4.5, 6.45, 4.3, 0.55, rounded=True, radius=0.5)
    solid(pill, ACCENT, alpha=12); line(pill, ACCENT, 1.0)
    pt = text(s, 4.5, 6.57, 4.3, 0.35, "Backend untouched · Zero database changes",
              size=12, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

    footer(s, 10)
    a.add(eb, "rise", 0.0, 0.5); a.add(title, "rise", 0.05, 0.5); a.add(sub, "rise", 0.2, 0.5)
    for i, grp in enumerate(rows):
        st = 0.5 + i * 0.1
        for sh in grp: a.add(sh, "rise", st, 0.35)
    a.add(pill, "rise", 1.5, 0.5); a.add(pt, "fade", 1.6, 0.4)
    apply_anim(s, a)
    transition(s, "morph", 800)
    return s

# ============================================================================
# SLIDE 11 - ROADMAP
# ============================================================================
def slide11():
    s = add_slide(); a = Anim()
    bg(s, BG_DEEP)
    eb = eyebrow(s, "10 · ROADMAP")
    title = text(s, 0.6, 1.0, 12, 0.9, "What comes next", size=44, color=TEXT, bold=True)
    sub = text(s, 0.6, 1.7, 12, 0.4, "Eight follow-up pull requests the team will tackle next.",
               size=14, color=TEXT_DIM)

    cards_data = [
        ("Login + Register views", "Marco and Rebeca, themed forms wired to Melissa's auth backend"),
        ("Events browse page", "Full list with the search and filter UI already drafted"),
        ("Event details page", "Single event view with the share modal trigger"),
        ("Create event form", "Posted to Joseph's EventsController.Create endpoint"),
        ("Edit event form", "Owners only, validated on the backend"),
        ("Delete confirmation", "Confirm modal with backend cascade handled"),
        ("My events dashboard", "Authorized view, lists what the user created or saved"),
        ("Share modal", "Access code copy-to-clipboard built on Melissa's share logic"),
    ]
    cw, ch = 3.0, 2.0
    gx = 0.13; gy = 0.2
    x0 = 0.6; y0 = 2.4
    accents = [ACCENT, VIOLET, PINK]
    groups = []
    for i, (ttl, desc) in enumerate(cards_data):
        col = i % 4; row = i // 4
        x = x0 + col * (cw + gx); y = y0 + row * (ch + gy)
        cd = card(s, x, y, cw, ch)
        col_c = accents[i % 3]
        num = text(s, x + 0.3, y + 0.22, 1.5, 0.6, f"{i+1:02d}", size=32, color=col_c, bold=True)
        h = text(s, x + 0.3, y + 0.85, cw - 0.6, 0.5, ttl, size=14, color=TEXT, bold=True, line_spacing=1.0)
        d = text(s, x + 0.3, y + 1.3, cw - 0.6, 0.6, desc, size=11, color=TEXT_DIM, line_spacing=1.15)
        groups.append([cd, num, h, d])

    footer(s, 11)
    a.add(eb, "rise", 0.0, 0.5); a.add(title, "rise", 0.05, 0.5); a.add(sub, "rise", 0.2, 0.5)
    for i, grp in enumerate(groups):
        st = 0.5 + i * 0.08
        for sh in grp: a.add(sh, "rise", st, 0.4)
    apply_anim(s, a)
    transition(s, "push", 700)
    return s

# ============================================================================
# SLIDE 12 - CLOSING (full-bleed G1)
# ============================================================================
def slide12():
    s = add_slide(); a = Anim()
    base = rect(s, -0.06, -0.06, SW + 0.12, SH + 0.12)
    gradient(base, [(0, ACCENT), (0.55, VIOLET), (1, PINK)], 135); no_line(base)

    head = text(s, 0.6, 2.55, 12.1, 1.4, "Frontend is ready.", size=96, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
    sub = text(s, 0.6, 3.95, 12.1, 0.8, "Backend, take it from here.", size=40, color="FFFFFF", align=PP_ALIGN.CENTER, alpha=92)
    body = text(s, 1.5, 4.95, 10.3, 0.6,
                "Pull the branch. Run it locally. Drop comments on the PR. Questions go to Marco.",
                size=18, color="FFFFFF", align=PP_ALIGN.CENTER, alpha=78)

    ll = text(s, 0.6, 6.75, 4, 0.3, "BRANCH", size=9, color="FFFFFF", bold=True, spc=3.0, caps=True, alpha=72)
    lv = text(s, 0.6, 7.05, 5, 0.3, "marco/frontend-design", size=14, color="FFFFFF", bold=True)
    rl = text(s, 7.7, 6.75, 5.0, 0.3, "TEAM", size=9, color="FFFFFF", bold=True, spc=3.0, caps=True, alpha=72, align=PP_ALIGN.RIGHT)
    rv = text(s, 7.7, 7.05, 5.0, 0.3, "The Syntax Errors · CSE 325", size=14, color="FFFFFF", bold=True, align=PP_ALIGN.RIGHT)

    a.add(head, "rise", 0.2, 0.8)
    a.add(sub, "rise", 0.9, 0.6)
    a.add(body, "rise", 1.5, 0.5)
    a.group([ll, lv, rl, rv], "rise", start=2.0, stagger=0.05, dur=0.4)
    apply_anim(s, a)
    transition(s, "morph", 800)
    return s

# ----------------------------------------------------------------------------
for fn in (slide1, slide2, slide3, slide4, slide5, slide6,
           slide7, slide8, slide9, slide10, slide11, slide12):
    fn()

out = "/tmp/output/notespack_frontend_v3.pptx"
prs.save(out)
print("saved", out)
