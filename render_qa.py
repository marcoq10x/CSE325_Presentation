#!/usr/bin/env python3
"""Rasterize the generated pptx to PNGs for visual QA (reads the real file)."""
import sys
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn

SCALE = 150  # px per inch
REG = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
BLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
DJV = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
_fc = {}
def font(size_pt, bold, special=False):
    px = max(6, int(size_pt * SCALE / 72))
    key = (px, bold, special)
    if key not in _fc:
        path = DJV if special else (BLD if bold else REG)
        _fc[key] = ImageFont.truetype(path, px)
    return _fc[key]
def _special(t):
    return any(ord(c) > 0x2000 for c in t)

def emu_px(v): return int(Emu(v).inches * SCALE)

def blend(hex_, alpha, bg=(7,7,26)):
    r=int(hex_[0:2],16); g=int(hex_[2:4],16); b=int(hex_[4:6],16)
    if alpha is None: return (r,g,b)
    a=alpha/100.0
    return (int(r*a+bg[0]*(1-a)), int(g*a+bg[1]*(1-a)), int(b*a+bg[2]*(1-a)))

def parse_fill(spPr):
    if spPr is None: return None
    sf = spPr.find(qn("a:solidFill"))
    if sf is not None:
        c = sf.find(qn("a:srgbClr"))
        al = c.find(qn("a:alpha"))
        return ("solid", c.get("val"), int(al.get("val"))/1000 if al is not None else None)
    gf = spPr.find(qn("a:gradFill"))
    if gf is not None:
        stops=[]
        for gs in gf.findall(qn("a:gsLst")+"/"+qn("a:gs")):
            c=gs.find(qn("a:srgbClr")); al=c.find(qn("a:alpha"))
            stops.append((int(gs.get("pos"))/100000, c.get("val"),
                          int(al.get("val"))/1000 if al is not None else None))
        radial = gf.find(qn("a:path")) is not None
        ang=0
        lin=gf.find(qn("a:lin"))
        if lin is not None: ang=int(lin.get("ang"))/60000
        return ("grad", stops, ang, radial)
    if spPr.find(qn("a:noFill")) is not None:
        return ("none",)
    return None

def grad_img(w,h,stops,ang,radial):
    img=Image.new("RGB",(max(1,w),max(1,h)))
    px=img.load()
    c0=blend(stops[0][1],stops[0][2]); c1=blend(stops[-1][1],stops[-1][2])
    import math
    if radial:
        cx,cy=w/2,h/2; mx=math.hypot(cx,cy)
        for y in range(h):
            for x in range(w):
                t=min(1,math.hypot(x-cx,y-cy)/mx)
                px[x,y]=tuple(int(c0[i]+(c1[i]-c0[i])*t) for i in range(3))
    else:
        rad=math.radians(ang)
        dx,dy=math.cos(rad),math.sin(rad)
        corners=[(0,0),(w,0),(0,h),(w,h)]
        projs=[cx*dx+cy*dy for cx,cy in corners]
        pmin,pmax=min(projs),max(projs); rng=(pmax-pmin) or 1
        stop_cols=[blend(s[1],s[2]) for s in stops]
        stop_pos=[s[0] for s in stops]
        for y in range(h):
            for x in range(w):
                t=((x*dx+y*dy)-pmin)/rng
                t=min(1,max(0,t))
                # interpolate across all stops
                for si in range(len(stop_pos)-1):
                    if t<=stop_pos[si+1] or si==len(stop_pos)-2:
                        lo,hi=stop_pos[si],stop_pos[si+1]
                        lt=(t-lo)/((hi-lo) or 1); lt=min(1,max(0,lt))
                        a,b=stop_cols[si],stop_cols[si+1]
                        px[x,y]=tuple(int(a[i]+(b[i]-a[i])*lt) for i in range(3))
                        break
    return img

def parse_line(spPr):
    ln=spPr.find(qn("a:ln"))
    if ln is None: return None
    nf=ln.find(qn("a:noFill"))
    if nf is not None: return None
    sf=ln.find(qn("a:solidFill"))
    if sf is None: return None
    c=sf.find(qn("a:srgbClr")); al=c.find(qn("a:alpha"))
    w=ln.get("w"); wpx=max(1,int(int(w)/12700*SCALE/72)) if w else 1
    return (blend(c.get("val"), int(al.get("val"))/1000 if al is not None else None), wpx)

def render_slide(slide, idx):
    W=int(13.333*SCALE); H=int(7.5*SCALE)
    img=Image.new("RGB",(W,H),(7,7,26))
    d=ImageDraw.Draw(img,"RGBA")
    for sh in slide.shapes:
        try: spPr=sh._element.spPr
        except Exception: spPr=None
        x=emu_px(sh.left or 0); y=emu_px(sh.top or 0)
        w=emu_px(sh.width or 0); h=emu_px(sh.height or 0)
        prst=None
        if spPr is not None:
            g=spPr.find(qn("a:prstGeom"))
            if g is not None: prst=g.get("prst")
        fill=parse_fill(spPr) if spPr is not None else None
        lin=parse_line(spPr) if spPr is not None else None
        rad=int(min(w,h)*0.12)
        if fill and fill[0]!="none":
            if fill[0]=="solid":
                if fill[2] is not None:  # alpha -> composite over actual pixels
                    r=int(fill[1][0:2],16); g=int(fill[1][2:4],16); b=int(fill[1][4:6],16)
                    col=(r,g,b,int(fill[2]*2.55))
                else:
                    col=blend(fill[1],fill[2])
                if prst=="ellipse": d.ellipse([x,y,x+w,y+h],fill=col)
                elif prst=="roundRect": d.rounded_rectangle([x,y,x+w,y+h],radius=rad,fill=col)
                else: d.rectangle([x,y,x+w,y+h],fill=col)
            elif fill[0]=="grad" and w>0 and h>0:
                gi=grad_img(w,h,fill[1],fill[2],fill[3])
                if prst=="roundRect":
                    mask=Image.new("L",(w,h),0); md=ImageDraw.Draw(mask)
                    md.rounded_rectangle([0,0,w-1,h-1],radius=rad,fill=255)
                    img.paste(gi,(x,y),mask)
                elif prst=="ellipse":
                    mask=Image.new("L",(w,h),0); md=ImageDraw.Draw(mask)
                    md.ellipse([0,0,w-1,h-1],fill=255); img.paste(gi,(x,y),mask)
                else: img.paste(gi,(x,y))
        if lin:
            col,wd=lin
            if prst=="ellipse": d.ellipse([x,y,x+w,y+h],outline=col,width=wd)
            elif prst=="roundRect": d.rounded_rectangle([x,y,x+w,y+h],radius=rad,outline=col,width=wd)
            else: d.rectangle([x,y,x+w,y+h],outline=col,width=wd)
        # text (with word wrapping to match PowerPoint word_wrap)
        if sh.has_text_frame and sh.text_frame.text.strip():
            from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
            tf=sh.text_frame
            anchor=tf.vertical_anchor
            avail=max(10, w-4)
            wrapped=[]  # list of (align, [ (txt,sz,bd,col, width) ])
            for p in tf.paragraphs:
                runs=[(r.text,(r.font.size.pt if r.font.size else 14),
                       bool(r.font.bold),_runcolor(r)) for r in p.runs if r.text]
                if not runs: continue
                al=p.alignment
                # flatten to words carrying style
                words=[]
                for txt,sz,bd,col in runs:
                    parts=txt.split(" ")
                    for k,wd in enumerate(parts):
                        token=wd+(" " if k<len(parts)-1 else "")
                        if token: words.append((token,sz,bd,col))
                line=[]; lw=0
                for token,sz,bd,col in words:
                    f=font(sz,bd,_special(token)); tl=d.textlength(token,font=f)
                    if line and lw+tl>avail:
                        wrapped.append((al,line)); line=[]; lw=0
                    line.append((token,sz,bd,col,tl)); lw+=tl
                if line: wrapped.append((al,line))
            lh=[int(max([t[1] for t in ln],default=14)*SCALE/72*1.28) for _,ln in wrapped]
            total=sum(lh)
            ty=y
            if anchor==MSO_ANCHOR.MIDDLE: ty=y+(h-total)//2
            elif anchor==MSO_ANCHOR.BOTTOM: ty=y+h-total
            cy=ty
            for (al,ln),lhh in zip(wrapped,lh):
                tw=sum(t[4] for t in ln)
                if al==PP_ALIGN.CENTER: cx=x+(w-tw)//2
                elif al==PP_ALIGN.RIGHT: cx=x+w-tw
                else: cx=x
                for token,sz,bd,col,ww in ln:
                    d.text((cx,cy),token,font=font(sz,bd,_special(token)),fill=col); cx+=ww
                cy+=lhh
    img.save(f"/tmp/qa/slide_{idx:02d}.png")
    return img

def _runcolor(r):
    try:
        rPr=r._r.find(qn("a:rPr"))
        if rPr is not None:
            sf=rPr.find(qn("a:solidFill"))
            if sf is not None:
                c=sf.find(qn("a:srgbClr")); al=c.find(qn("a:alpha"))
                return blend(c.get("val"), int(al.get("val"))/1000 if al is not None else None)
    except Exception: pass
    return (234,241,255)

def main():
    path=sys.argv[1] if len(sys.argv)>1 else "/tmp/output/notespack_frontend_v3.pptx"
    prs=Presentation(path)
    imgs=[]
    for i,s in enumerate(prs.slides,1):
        imgs.append(render_slide(s,i))
    # contact sheet 3x4
    cols,rows=3,4; tw=imgs[0].width//3; th=imgs[0].height//3
    sheet=Image.new("RGB",(tw*cols+40, th*rows+50),(20,20,30))
    for i,im in enumerate(imgs):
        t=im.resize((tw,th))
        c=i%cols; r=i//cols
        sheet.paste(t,(10+c*tw+c*5, 10+r*th+r*5))
    sheet.save("/tmp/qa/contact_sheet.png")
    print("rendered",len(imgs),"slides + contact sheet")

if __name__=="__main__":
    main()
